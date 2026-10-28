"""Merge the supplied national food workbook with prefectural TFR by year.

Reads cached spreadsheet values only; workbook edits are made by the JS builder.
"""
import argparse
import json
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("payload", type=Path)
    args = parser.parse_args()
    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    tfr = pd.read_csv(ROOT / "data/prefectural_tfr_2016_2025.csv")
    names = pd.read_csv(ROOT / "data/dated_source_panel_2016_2025.csv",
                        usecols=["region_code", "prefecture", "prefecture_japanese"]).drop_duplicates()
    assert names.region_code.is_unique
    tfr = tfr.merge(names, on="region_code", how="left", validate="many_to_one")
    assert not tfr.duplicated(["region_code", "year"]).any() and tfr.prefecture.notna().all()
    tfr = tfr.sort_values(["region_code", "year"]).reset_index(drop=True)
    annual, soy = wb["Annual groups"], wb["Soyabeans"]
    years = [annual.cell(5, c).value for c in range(4, 14)]
    assert years == list(range(2016, 2026))
    codes = [(str(annual.cell(r, 2).value).zfill(2), annual.cell(r, 3).value, r)
             for r in range(6, 27)]
    assert len(set(c[0] for c in codes)) == 21
    soy_codes = [(str(soy.cell(r, 2).value).zfill(6), soy.cell(r, 3).value, r)
                 for r in range(6, 13)]
    fields = ["year", "published_food_usd",
              "mapped_food_usd", "whole_soyabeans_usd", "other_mapped_food_usd"]
    fields += ["soy_hs"+c+"_usd" for c, _, _ in soy_codes]
    fields += ["group_hs"+c+"_usd" for c, _, _ in codes]
    rows = []
    for j, year in enumerate(years):
        col = 4+j
        def value(x):
            if x in [None, "n.a."]:
                return None
            if not isinstance(x, (int, float)):
                raise ValueError(f"Unexpected import value: {x}")
            return x
        soy_values = [value(soy.cell(r, 5+j).value) for _, _, r in soy_codes]
        mapped = value(annual.cell(28, col).value)
        published = value(annual.cell(29, col).value)
        whole = soy_values[0] + soy_values[1]
        rows.append([year, published, mapped, whole, mapped-whole] + soy_values +
                    [value(annual.cell(r, col).value) for _, _, r in codes])
    imports = pd.DataFrame(rows, columns=fields)
    assert imports.year.is_unique
    merged = tfr.merge(imports, on="year", how="left", validate="many_to_one", indicator=True)
    assert merged._merge.eq("both").all() and merged.tfr_children_per_woman.notna().all()
    merged["import_geographic_level"] = "Japan national; shared within year; not prefectural exposure"
    fields = ["region_code", "prefecture", "prefecture_japanese", "year", "tfr_children_per_woman",
              "tfr_release_status", "import_geographic_level"] + fields[1:]
    merged = merged[fields + ["tfr_source", "tfr_source_cell"]]
    assert len(merged) == 470 and merged.region_code.nunique() == 47
    assert merged.groupby("year").size().eq(47).all()
    dest = ROOT / "output/tfr_merged"
    dest.mkdir(parents=True, exist_ok=True)
    merged.to_csv(dest/"prefecture_food_imports_and_tfr_2016_2025.csv", index=False)
    merged[merged.year.le(2024)].to_csv(dest/"prefecture_food_imports_and_tfr_final_2016_2024.csv", index=False)
    payload = dict(tfr=tfr.to_dict("records"), fields=fields, rows=merged[fields].astype(object).where(pd.notna(merged[fields]), None).values.tolist(),
                   group_codes=codes, soy_codes=soy_codes)
    args.payload.parent.mkdir(parents=True, exist_ok=True)
    args.payload.write_text(json.dumps(payload, allow_nan=False))
    print(f"Merged {len(merged)} prefecture-years; {len(fields)} fields; no unmatched years or duplicate prefecture-year keys.")


if __name__ == "__main__":
    main()
